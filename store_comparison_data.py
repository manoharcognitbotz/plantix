import os
import openpyxl
import psycopg2
from psycopg2.extensions import ISOLATION_LEVEL_AUTOCOMMIT
from dotenv import load_dotenv

# Absolute path to the project directory for RPA compatibility
try:
    PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
except NameError:
    PROJECT_DIR = r"C:\UiPathProjects\BlankProcess\SE"
ENV_PATH = os.path.join(PROJECT_DIR, ".env")

def create_database_if_not_exists():
    # Load configuration
    load_dotenv(ENV_PATH)
    db_host = os.getenv("DB_HOST")
    db_port = os.getenv("DB_PORT")
    db_user = os.getenv("DB_USER")
    db_password = os.getenv("DB_PASSWORD")
    db_name = os.getenv("DB_NAME")

    # Connect to default database 'postgres' to check/create the database
    print(f"[*] Connecting to database server {db_host}:{db_port} as user '{db_user}'...")
    conn = None
    cursor = None
    try:
        conn_params = {
            "host": db_host,
            "port": db_port,
            "user": db_user,
            "password": db_password,
            "database": "postgres"
        }
        if db_host and "neon.tech" in db_host:
            conn_params["sslmode"] = "require"
        conn = psycopg2.connect(**conn_params)
        conn.set_isolation_level(ISOLATION_LEVEL_AUTOCOMMIT)
        cursor = conn.cursor()

        # Check if database exists
        cursor.execute("SELECT 1 FROM pg_catalog.pg_database WHERE datname = %s;", (db_name,))
        exists = cursor.fetchone()
        
        if not exists:
            print(f"[*] Database '{db_name}' does not exist. Creating it...")
            cursor.execute(f'CREATE DATABASE "{db_name}";')
            print(f"[+] Database '{db_name}' created successfully.")
        else:
            print(f"[*] Database '{db_name}' already exists.")

    except Exception as e:
        print(f"[*] Skipping database creation check (typical for cloud-hosted databases like Neon): {e}")
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()

def import_excel_to_db():
    load_dotenv(ENV_PATH)
    db_host = os.getenv("DB_HOST")
    db_port = os.getenv("DB_PORT")
    db_user = os.getenv("DB_USER")
    db_password = os.getenv("DB_PASSWORD")
    db_name = os.getenv("DB_NAME")
    
    comparision_dir = os.path.join(PROJECT_DIR, "Comparision data")
    if not os.path.exists(comparision_dir):
        print(f"[-] Error: Comparision data directory not found at {comparision_dir}")
        return

    # Find all Excel files in the folder
    excel_files = []
    for f in os.listdir(comparision_dir):
        if f.startswith('~$'):
            continue
        if os.path.splitext(f)[1].lower() in ['.xlsx', '.xls']:
            excel_files.append(os.path.join(comparision_dir, f))

    if not excel_files:
        print(f"[-] No Excel files found in {comparision_dir}")
        return

    print(f"[*] Found {len(excel_files)} Excel file(s) to import.")
    print(f"[*] Connecting to database '{db_name}'...")
    
    conn = None
    cursor = None
    try:
        conn_params = {
            "host": db_host,
            "port": db_port,
            "user": db_user,
            "password": db_password,
            "database": db_name
        }
        if db_host and "neon.tech" in db_host:
            conn_params["sslmode"] = "require"
        conn = psycopg2.connect(**conn_params)
        cursor = conn.cursor()

        # Create table only if it doesn't already exist
        print("[*] Checking table 'comparison_data'...")
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS comparison_data (
                id SERIAL PRIMARY KEY,
                pr_item_code TEXT NOT NULL,
                uom VARCHAR(50),
                indent_quantity NUMERIC,
                imported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
        """)
        conn.commit()

        # Check if table already has data and we are in the middle of a batch
        cursor.execute("SELECT COUNT(*) FROM comparison_data;")
        db_count = cursor.fetchone()[0]
        
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables 
                WHERE table_schema = 'public' 
                AND table_name = 'processed_files'
            );
        """)
        pf_exists = cursor.fetchone()[0]
        
        processed_count = 0
        if pf_exists:
            cursor.execute("SELECT COUNT(*) FROM processed_files;")
            processed_count = cursor.fetchone()[0]
            
        if db_count > 0 and processed_count > 0:
            print("[*] Active batch detected. Preserving comparison_data and existing supplier items.")
            return

        # Truncate comparison_data table to do a clean import (prevent duplicates)
        print("[*] Clearing table 'comparison_data' for clean import...")
        cursor.execute("TRUNCATE comparison_data CASCADE;")
        conn.commit()

        insert_query = """
            INSERT INTO comparison_data (pr_item_code, uom, indent_quantity)
            VALUES (%s, %s, %s);
        """

        total_inserted = 0
        for excel_path in excel_files:
            filename = os.path.basename(excel_path)
            print(f"[*] Reading data from '{filename}'...")
            try:
                wb = openpyxl.load_workbook(excel_path, data_only=True)
                sheet = wb.active  # Load the active sheet
                
                rows = list(sheet.iter_rows(values_only=True))
                if len(rows) <= 1:
                    print(f"[-] Warning: Excel sheet in '{filename}' is empty or contains only headers.")
                    continue

                data_rows = rows[1:]
                inserted_count = 0
                for row in data_rows:
                    if not row or len(row) < 3:
                        continue
                    pr_item_code = row[0]
                    uom = row[1]
                    indent_quantity = row[2]

                    # Skip rows where pr_item_code is null/empty
                    if pr_item_code is None:
                        continue
                    
                    # Strip string values if they are text
                    if isinstance(pr_item_code, str):
                        pr_item_code = pr_item_code.strip()
                    if isinstance(uom, str):
                        uom = uom.strip()

                    cursor.execute(insert_query, (pr_item_code, uom, indent_quantity))
                    inserted_count += 1
                
                conn.commit()
                print(f"[+] Successfully imported {inserted_count} rows from '{filename}'.")
                total_inserted += inserted_count
            except Exception as e:
                print(f"[-] Failed to import from '{filename}': {e}")

        print(f"[+] Import complete. Total rows inserted: {total_inserted}")

        # Verify insertion
        cursor.execute("SELECT id, pr_item_code, uom, indent_quantity FROM comparison_data ORDER BY id;")
        db_rows = cursor.fetchall()
        print(f"\n--- Verification: Contents of table 'comparison_data' ({len(db_rows)} records) ---")
        for r in db_rows:
            print(f"ID: {r[0]} | Item: {r[1]} | UOM: {r[2]} | Qty: {r[3]}")

    except Exception as e:
        print(f"[-] Error during Excel database import: {e}")
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()

if __name__ == "__main__":
    create_database_if_not_exists()
    import_excel_to_db()
